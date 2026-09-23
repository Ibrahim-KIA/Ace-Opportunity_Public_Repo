"""File text extraction for PDF and DOCX CV documents."""
from __future__ import annotations

import io
import os
from typing import Set

ALLOWED_EXTENSIONS: Set[str] = {".pdf", ".docx"}
MAX_FILE_BYTES: int = 5 * 1024 * 1024  # 5 MB


def extract_text(filename: str, file_bytes: bytes, content_type: str = "") -> str:
    """Extract plain text from a PDF or DOCX file bytes.

    Args:
        filename: Name of the uploaded file.
        file_bytes: Raw bytes of the file.
        content_type: Optional MIME content type.

    Returns:
        Extracted plain text string.

    Raises:
        ValueError: if the file type is unsupported, file is empty, or text cannot be extracted.
    """
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    ext = get_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF or DOCX document."
        )

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")


def get_extension(filename: str) -> str:
    """Return lowercase file extension with dot, e.g. '.pdf'."""
    if not filename:
        return ""
    _, ext = os.path.splitext(filename.lower())
    return ext


def _extract_pdf(data: bytes) -> str:
    """Extract text from PDF bytes using pypdf (with pdfplumber fallback)."""
    text_content = ""
    # Try pypdf first
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            p_text = page.extract_text()
            if p_text:
                pages.append(p_text.strip())
        text_content = "\n".join(pages).strip()
    except Exception as exc:
        # Fallback to pdfplumber if available
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = [p.extract_text() or "" for p in pdf.pages]
            text_content = "\n".join(pages).strip()
        except ImportError:
            raise ValueError(f"Failed to parse PDF document: {exc}") from exc

    if not text_content:
        raise ValueError(
            "Could not extract any text from the PDF. Is it a scanned image or empty document?"
        )
    return text_content


def _extract_docx(data: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also extract table cells if any
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    ctext = cell.text.strip()
                    if ctext and ctext not in table_texts:
                        table_texts.append(ctext)

        combined = paragraphs + table_texts
        text = "\n".join(combined).strip()
    except Exception as exc:
        raise ValueError(f"Failed to parse DOCX document: {exc}") from exc

    if not text:
        raise ValueError(
            "Could not extract any text from the DOCX file. The document appears empty."
        )
    return text
